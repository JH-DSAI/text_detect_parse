"""Document parsing for txt, pdf and docx files."""

from __future__ import annotations

import re
import warnings
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .models import Block, Document, ImageBlock, TableBlock, TextBlock

if TYPE_CHECKING:
    import pymupdf
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

_Element = Any  # lxml ships no type stubs

SUPPORTED_EXTENSIONS = (".txt", ".pdf", ".docx")


def parse(
    path: str | Path,
    *,
    parse_images: bool = True,
    parse_tables: bool = True,
) -> Document:
    """Parse a document into an ordered list of text, table and image blocks.

    When `parse_tables` / `parse_images` is False that content is skipped
    entirely (table text is not leaked into the text blocks either).
    Plain-text files contain neither, so the flags have no effect on them.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    try:
        parser = _PARSERS[path.suffix.lower()]
    except KeyError:
        raise ValueError(
            f"Unsupported file type {path.suffix!r}; supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        ) from None
    return Document(source=path, blocks=parser(path, parse_images, parse_tables))


def _parse_txt(path: Path, parse_images: bool, parse_tables: bool) -> list[Block]:
    try:
        content = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        content = path.read_text(encoding="latin-1")
    paragraphs = re.split(r"\n\s*\n", content.replace("\r\n", "\n"))
    return [TextBlock(p.strip()) for p in paragraphs if p.strip()]


def _parse_docx(path: Path, parse_images: bool, parse_tables: bool) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    blocks: list[Block] = []
    blip_embed = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    embed_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

    for child in _iter_body(doc.element.body):
        tag = _local_name(child)
        if tag == "p":
            para = Paragraph(child, doc)
            level = _heading_level(para)
            pending: list[str] = []

            def flush_text(level: int = level, pending: list[str] = pending) -> None:
                text = "".join(pending).strip()
                pending.clear()
                if text:
                    blocks.append(TextBlock(text, level=level))

            # Walk runs in order so text after an inline image stays after it.
            for item in para.iter_inner_content():
                for run in item.runs if isinstance(item, Hyperlink) else [item]:
                    pending.append(run.text)
                    if not parse_images:
                        continue
                    for blip in run._r.iter(blip_embed):
                        part = doc.part.related_parts.get(blip.get(embed_attr))
                        if part is not None:
                            flush_text()
                            ext = Path(part.partname).suffix.lstrip(".") or "png"
                            blocks.append(ImageBlock(part.blob, ext=ext))
            flush_text()
        elif tag == "tbl" and parse_tables:
            table = Table(child, doc)
            blocks.append(TableBlock(_docx_rows(table)))
    return blocks


def _local_name(element: _Element) -> str:
    return str(element.tag).rsplit("}", 1)[-1]


def _iter_body(parent: _Element) -> Iterator[_Element]:
    """Yield paragraph/table elements in order, unwrapping content controls (w:sdt)."""
    for child in parent.iterchildren():
        tag = _local_name(child)
        if tag == "sdt":
            for content in child.iterchildren():
                if _local_name(content) == "sdtContent":
                    yield from _iter_body(content)
        elif tag in ("p", "tbl"):
            yield child


def _heading_level(para: Paragraph) -> int:
    # Walk base styles so custom styles derived from a heading are recognised.
    style = para.style
    while style is not None:
        name = style.name or ""
        if m := re.fullmatch(r"Heading (\d)", name):
            return min(int(m.group(1)), 6)
        if name == "Title":
            return 1
        style = style.base_style
    return 0


def _docx_rows(table: Table) -> list[list[str]]:
    # A horizontally merged cell shows up several times in one row: keep it once.
    # A vertically merged cell shows up in every row it spans: keep the text in the
    # first row and leave "" below so columns stay aligned.
    seen: set[_Element] = set()
    rows = []
    for row in table.rows:
        in_row: set[_Element] = set()
        cells = []
        for cell in row.cells:
            if cell._tc in in_row:
                continue
            in_row.add(cell._tc)
            cells.append("" if cell._tc in seen else _docx_cell_text(cell))
            seen.add(cell._tc)
        rows.append(cells)
    return rows


def _docx_cell_text(cell: _Cell) -> str:
    # cell.text ignores nested tables, so walk the cell body ourselves.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parts = []
    for child in _iter_body(cell._tc):
        if _local_name(child) == "p":
            parts.append(Paragraph(child, cell).text)
        else:
            parts.extend(" | ".join(r) for r in _docx_rows(Table(child, cell)))
    return "\n".join(parts).strip()


def _parse_pdf(path: Path, parse_images: bool, parse_tables: bool) -> list[Block]:
    import pymupdf

    blocks: list[Block] = []
    with pymupdf.open(path) as pdf:
        for page_no, page in enumerate(pdf, 1):
            # Always locate tables so their text is excluded from the text
            # blocks, even when table parsing is disabled.
            try:
                tables = page.find_tables().tables
            except Exception as e:
                warnings.warn(f"Table detection failed on page {page_no}: {e}", stacklevel=2)
                tables = []
            table_boxes = [pymupdf.Rect(t.bbox) for t in tables]

            # (y0, block) in PyMuPDF's native order, which follows the content stream
            # and so keeps columns together; sorting by position would interleave them.
            found: list[tuple[float, Block]] = []
            for b in page.get_text("dict", sort=False)["blocks"]:
                box = pymupdf.Rect(b["bbox"])
                if b["type"] == 1:
                    if parse_images:
                        img = ImageBlock(b["image"], ext=b.get("ext", "png"), page=page_no)
                        found.append((box.y0, img))
                elif not any(_mostly_inside(box, t) for t in table_boxes):
                    lines = ("".join(s["text"] for s in ln["spans"]) for ln in b["lines"])
                    text = "\n".join(ln for ln in lines if ln.strip()).strip()
                    if text:
                        found.append((box.y0, TextBlock(text, page=page_no)))

            if parse_tables:
                # Tables go before the first block that starts at or below them.
                for t, box in sorted(
                    zip(tables, table_boxes, strict=True), key=lambda tb: tb[1].y0, reverse=True
                ):
                    rows = [[c or "" for c in r] for r in t.extract()]
                    at = next((i for i, (y, _) in enumerate(found) if y >= box.y0), len(found))
                    found.insert(at, (box.y0, TableBlock(rows, page=page_no)))

            blocks += [blk for _, blk in found]
    return blocks


def _mostly_inside(box: pymupdf.Rect, container: pymupdf.Rect) -> bool:
    area = box.get_area()
    return bool(area > 0 and (box & container).get_area() / area > 0.5)


_PARSERS: dict[str, Callable[[Path, bool, bool], list[Block]]] = {
    ".txt": _parse_txt,
    ".docx": _parse_docx,
    ".pdf": _parse_pdf,
}
