"""Shared fixtures: small synthetic txt/docx/pdf documents built in code."""

import io
from pathlib import Path

import docx
import pymupdf
import pytest


def make_pixmap(fmt: str = "png") -> bytes:
    return pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False).tobytes(fmt)


PNG = make_pixmap("png")


@pytest.fixture
def png_bytes() -> bytes:
    return PNG


@pytest.fixture
def txt_file(tmp_path: Path) -> Path:
    p = tmp_path / "a.txt"
    p.write_text("First para\nstill first\n\n\nSecond para\n")
    return p


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Title here", level=1)
    d.add_paragraph("Body text")
    t = d.add_table(rows=2, cols=2)
    for i, v in enumerate(["h1", "h2", "c1", "c2"]):
        t.cell(i // 2, i % 2).text = v
    d.add_picture(io.BytesIO(PNG))
    d.add_paragraph("After")
    p = tmp_path / "a.docx"
    d.save(str(p))
    return p


def draw_table(page: pymupdf.Page, x0: float, y0: float, rows: int = 3, cols: int = 2) -> None:
    w, h = 100, 30
    for r in range(rows):
        for c in range(cols):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect)
            page.insert_text((rect.x0 + 5, rect.y0 + 20), f"r{r}c{c}")


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Intro paragraph")
    draw_table(page, 72, 120)
    page.insert_image(pymupdf.Rect(72, 300, 172, 400), stream=PNG)
    page.insert_text((72, 450), "Outro paragraph")
    p = tmp_path / "a.pdf"
    doc.save(p)
    return p
