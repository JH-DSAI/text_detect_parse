import io
from pathlib import Path

import docx
import pymupdf
import pytest
from conftest import PNG, draw_table

from text_detect_parse import Document, ImageBlock, TableBlock, TextBlock, parse


def kinds(doc: Document) -> list[str]:
    return [b.type for b in doc.blocks]


def texts(doc: Document) -> list[TextBlock]:
    return [b for b in doc.blocks if isinstance(b, TextBlock)]


# ---- general ---------------------------------------------------------------


def test_unsupported_extension(tmp_path: Path) -> None:
    (tmp_path / "x.csv").write_text("a")
    with pytest.raises(ValueError, match="Unsupported"):
        parse(tmp_path / "x.csv")


def test_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        parse(tmp_path / "nope.txt")


def test_accepts_str_path_and_uppercase_extension(tmp_path: Path) -> None:
    p = tmp_path / "A.TXT"
    p.write_text("hello")
    assert parse(str(p)).text == "hello"


# ---- txt -------------------------------------------------------------------


def test_txt_paragraphs(txt_file: Path) -> None:
    doc = parse(txt_file, parse_images=False, parse_tables=False)
    assert [b.text for b in texts(doc)] == ["First para\nstill first", "Second para"]


def test_txt_flags_are_noops(txt_file: Path) -> None:
    assert parse(txt_file).blocks == parse(txt_file, parse_images=False, parse_tables=False).blocks


def test_txt_empty(tmp_path: Path) -> None:
    p = tmp_path / "e.txt"
    p.write_text("  \n\n  \n")
    assert parse(p).blocks == []


def test_txt_crlf_and_bom(tmp_path: Path) -> None:
    p = tmp_path / "w.txt"
    p.write_bytes(b"\xef\xbb\xbfone\r\n\r\ntwo\r\n")
    assert [b.text for b in texts(parse(p))] == ["one", "two"]


def test_txt_bare_cr_and_utf16(tmp_path: Path) -> None:
    p = tmp_path / "cr.txt"
    p.write_bytes(b"one\r\rtwo")
    assert [b.text for b in texts(parse(p))] == ["one", "two"]
    u = tmp_path / "u16.txt"
    u.write_text("caf\u00e9\n\nnew", encoding="utf-16")
    assert [b.text for b in texts(parse(u))] == ["caf\u00e9", "new"]


def test_txt_latin1_fallback(tmp_path: Path) -> None:
    p = tmp_path / "l.txt"
    p.write_bytes("caf\xe9".encode("latin-1"))
    assert parse(p).text == "café"


# ---- docx ------------------------------------------------------------------


def test_docx_all(docx_file: Path) -> None:
    doc = parse(docx_file)
    assert kinds(doc) == ["text", "text", "table", "image", "text"]
    assert texts(doc)[0].level == 1
    assert doc.tables[0].rows == [["h1", "h2"], ["c1", "c2"]]
    assert doc.images[0].data == PNG
    assert doc.images[0].ext == "png"


def test_docx_flags(docx_file: Path) -> None:
    assert "table" not in kinds(parse(docx_file, parse_tables=False))
    assert "image" not in kinds(parse(docx_file, parse_images=False))
    assert kinds(parse(docx_file, parse_images=False, parse_tables=False)) == ["text"] * 3


def test_docx_table_text_not_leaked_when_disabled(docx_file: Path) -> None:
    assert "c1" not in parse(docx_file, parse_tables=False).text


def test_docx_heading_levels_and_title(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("T", level=0)
    d.add_heading("H2", level=2)
    d.add_heading("H9", level=9)
    d.add_paragraph("body")
    p = tmp_path / "h.docx"
    d.save(str(p))
    assert [(b.text, b.level) for b in texts(parse(p))] == [
        ("T", 1),
        ("H2", 2),
        ("H9", 6),
        ("body", 0),
    ]


def test_docx_blank_paragraphs_skipped(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("")
    d.add_paragraph("   ")
    d.add_paragraph("x")
    p = tmp_path / "b.docx"
    d.save(str(p))
    assert [b.text for b in texts(parse(p))] == ["x"]


def test_docx_merged_cells_not_duplicated(tmp_path: Path) -> None:
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "merged"
    t.cell(1, 0).text = "a"
    t.cell(1, 1).text = "b"
    p = tmp_path / "m.docx"
    d.save(str(p))
    assert parse(p).tables[0].rows == [["merged"], ["a", "b"]]


def test_docx_vertical_merge_not_repeated(tmp_path: Path) -> None:
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(2, 0)).text = "tall"
    for i in range(3):
        t.cell(i, 1).text = str(i)
    p = tmp_path / "v.docx"
    d.save(str(p))
    assert parse(p).tables[0].rows == [["tall", "0"], ["", "1"], ["", "2"]]


def test_docx_nested_table_text_kept(tmp_path: Path) -> None:
    d = docx.Document()
    outer = d.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.text = "outer"
    inner = cell.add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "in1"
    inner.cell(0, 1).text = "in2"
    p = tmp_path / "n.docx"
    d.save(str(p))
    assert parse(p).tables[0].rows == [["outer\nin1 | in2"]]


def test_docx_content_control_body_not_dropped(tmp_path: Path) -> None:
    from docx.oxml import parse_xml

    d = docx.Document()
    para = d.add_paragraph("inside control")
    xml = (
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:sdtContent/></w:sdt>"
    )
    sdt = parse_xml(xml)
    para._p.addprevious(sdt)
    sdt[0].append(para._p)
    p = tmp_path / "s.docx"
    d.save(str(p))
    assert [b.text for b in texts(parse(p))] == ["inside control"]


def test_docx_custom_style_based_on_heading(tmp_path: Path) -> None:
    from docx.enum.style import WD_STYLE_TYPE

    d = docx.Document()
    style = d.styles.add_style("My Heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = d.styles["Heading 2"]
    d.add_paragraph("custom", style="My Heading")
    p = tmp_path / "c.docx"
    d.save(str(p))
    assert [(b.text, b.level) for b in texts(parse(p))] == [("custom", 2)]


def test_docx_multiple_images_in_order(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("before")
    d.add_picture(io.BytesIO(PNG))
    d.add_paragraph("middle")
    d.add_picture(io.BytesIO(PNG))
    p = tmp_path / "i.docx"
    d.save(str(p))
    assert kinds(parse(p)) == ["text", "image", "text", "image"]


def test_docx_inline_image_splits_paragraph_text(tmp_path: Path) -> None:
    d = docx.Document()
    para = d.add_paragraph("before ")
    para.add_run().add_picture(io.BytesIO(PNG))
    para.add_run(" after")
    p = tmp_path / "inline.docx"
    d.save(str(p))
    doc = parse(p)
    assert kinds(doc) == ["text", "image", "text"]
    assert [b.text for b in texts(doc)] == ["before", "after"]
    assert [b.text for b in texts(parse(p, parse_images=False))] == ["before  after"]


def _docx_with_paragraph_xml(tmp_path: Path, inner: str) -> Path:
    from docx.oxml import parse_xml

    d = docx.Document()
    para = d.add_paragraph()
    rid, _ = d.part.get_or_add_image(io.BytesIO(PNG))
    ns = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    )
    for el in parse_xml(f"<w:p {ns}>{inner.replace('RID', rid)}</w:p>"):
        para._p.append(el)
    p = tmp_path / "x.docx"
    d.save(str(p))
    return p


def test_docx_image_between_text_in_one_run(tmp_path: Path) -> None:
    p = _docx_with_paragraph_xml(
        tmp_path,
        '<w:r><w:t>A</w:t><w:drawing><a:blip r:embed="RID"/></w:drawing><w:t>B</w:t></w:r>',
    )
    doc = parse(p)
    assert kinds(doc) == ["text", "image", "text"]
    assert [b.text for b in texts(doc)] == ["A", "B"]


def test_docx_vml_image_and_alternate_content_not_duplicated(tmp_path: Path) -> None:
    p = _docx_with_paragraph_xml(
        tmp_path,
        '<w:r><w:pict><v:shape><v:imagedata r:id="RID"/></v:shape></w:pict></w:r>'
        "<w:r><mc:AlternateContent>"
        '<mc:Choice><w:drawing><a:blip r:embed="RID"/></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:imagedata r:id="RID"/></w:pict></mc:Fallback>'
        "</mc:AlternateContent></w:r>",
    )
    assert kinds(parse(p)) == ["image", "image"]
    assert kinds(parse(p, parse_images=False)) == []


def test_docx_inline_wrappers_text_kept(tmp_path: Path) -> None:
    p = _docx_with_paragraph_xml(
        tmp_path,
        "<w:r><w:t>a </w:t></w:r>"
        "<w:sdt><w:sdtContent><w:r><w:t>b </w:t></w:r></w:sdtContent></w:sdt>"
        "<w:ins><w:r><w:t>c</w:t></w:r></w:ins>",
    )
    assert [b.text for b in texts(parse(p))] == ["a b c"]


def test_docx_corrupt(tmp_path: Path) -> None:
    p = tmp_path / "bad.docx"
    p.write_bytes(b"not a zip")
    with pytest.raises(Exception, match=r".+"):
        parse(p)


# ---- pdf -------------------------------------------------------------------


def test_pdf_all(pdf_file: Path) -> None:
    doc = parse(pdf_file)
    assert kinds(doc) == ["text", "table", "image", "text"]
    assert doc.tables[0].rows == [["r0c0", "r0c1"], ["r1c0", "r1c1"], ["r2c0", "r2c1"]]
    assert doc.tables[0].page == 1
    assert isinstance(doc.blocks[2], ImageBlock) and doc.blocks[2].data


def test_pdf_flags_skip_content(pdf_file: Path) -> None:
    doc = parse(pdf_file, parse_tables=False, parse_images=False)
    assert kinds(doc) == ["text", "text"]
    assert "r0c0" not in doc.text


def test_pdf_tables_only_and_images_only(pdf_file: Path) -> None:
    assert kinds(parse(pdf_file, parse_images=False)) == ["text", "table", "text"]
    assert kinds(parse(pdf_file, parse_tables=False)) == ["text", "image", "text"]


def test_pdf_multipage_numbers_and_order(tmp_path: Path) -> None:
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "page one")
    p2 = pdf.new_page()
    p2.insert_text((72, 72), "page two")
    draw_table(p2, 72, 120, rows=2)
    path = tmp_path / "m.pdf"
    pdf.save(path)
    doc = parse(path)
    assert [(b.type, b.page) for b in doc.blocks] == [("text", 1), ("text", 2), ("table", 2)]
    assert all(isinstance(b, TableBlock) for b in doc.tables)


def test_pdf_empty_page(tmp_path: Path) -> None:
    pdf = pymupdf.open()
    pdf.new_page()
    path = tmp_path / "e.pdf"
    pdf.save(path)
    assert parse(path).blocks == []


def test_pdf_corrupt(tmp_path: Path) -> None:
    p = tmp_path / "bad.pdf"
    p.write_bytes(b"not a pdf")
    with pytest.raises(Exception, match=r".+"):
        parse(p)


def test_pdf_columns_stay_together(tmp_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    for y, label in [(100, "L1"), (300, "L2")]:
        page.insert_text((72, y), label)
    for y, label in [(101, "R1"), (301, "R2")]:
        page.insert_text((350, y), label)
    p = tmp_path / "cols.pdf"
    doc.save(p)
    assert [b.text for b in texts(parse(p))] == ["L1", "L2", "R1", "R2"]


def test_pdf_table_placed_within_its_column(tmp_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    # Content-stream order: left column first, then right column above the table.
    page.insert_text((72, 100), "L1")
    page.insert_text((72, 700), "L2")
    page.insert_text((350, 100), "R1")
    draw_table(page, 350, 200, rows=2, cols=2)
    page.insert_text((350, 400), "R2")
    p = tmp_path / "t.pdf"
    doc.save(p)
    out = parse(p)
    assert [b.text if isinstance(b, TextBlock) else "T" for b in out.blocks] == [
        "L1",
        "L2",
        "R1",
        "T",
        "R2",
    ]


def test_pdf_table_detection_failure_warns_and_continues(
    pdf_file: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def boom(self: pymupdf.Page, *a: object, **k: object) -> None:
        raise RuntimeError("boom")

    monkeypatch.setattr(pymupdf.Page, "find_tables", boom)
    with pytest.warns(UserWarning, match="Table detection failed") as rec:
        doc = parse(pdf_file, parse_tables=False)
    assert rec[0].filename == __file__
    assert "Intro paragraph" in doc.text
